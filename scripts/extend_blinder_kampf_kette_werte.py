"""LLM-001 (SDF-013): Blinder-Kampf-Vorteilskette von 12 auf 18 Stufen erweitern.

Owner-Entscheidung 2026-09-16: Blinder Kampf ist in dieser SPOT keine Sonderfertigkeit mit
Maximum-Wert (wie im nasus-nasus-Ticket angenommen), sondern eine reine V/N-Vorteilskette
(vn_sicht_blinder_kampf_i.._12, RC-014s Max-Talent-Mechanik hat hier kein Ziel). Umsetzung
daher als direkte Kettenverlaengerung statt als Talent: 6 neue Stufen 13-18, exakt dasselbe
Muster wie die bestehenden Stufen 2-12 (25 SP, VORAUSSETZUNG auf die Vorstufe,
BLINDER_KAMPF_STUFE-Flag, "reduziert um insgesamt N"-Wirkungstext). Ein voll investierter
Charakter kann die permanente Sichtbehinderung 18 (vn_sicht_blindheit, DEC-220) damit
vollstaendig kompensieren.

Aktualisiert zusaetzlich den Info-Text aller 12 bestehenden Stufen (inkl. Basisstufe I,
Zeile 1107) von "Zwoelfstufige" auf "Achtzehnstufige Vorteilskette".
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

XLSX_PATH = Path(r"E:\Das Western Rollenspiel\LLM\werte 0.8-claude.xlsx")

ROEMISCH = {
    1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII", 8: "VIII", 9: "IX",
    10: "X", 11: "XI", 12: "XII", 13: "XIII", 14: "XIV", 15: "XV", 16: "XVI", 17: "XVII", 18: "XVIII",
}


def main():
    backup_path = XLSX_PATH.with_name(
        f"werte 0.8-claude_backup_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    )
    shutil.copy2(XLSX_PATH, backup_path)
    print(f"Backup angelegt: {backup_path}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    ws = wb["Werte"]

    # 1) Info-Text der 12 bestehenden Stufen aktualisieren ("Zwoelfstufige" -> "Achtzehnstufige").
    updated = 0
    for r in range(1, ws.max_row + 1):
        ref = ws.cell(row=r, column=1).value
        if not (ref and str(ref).startswith("vn_sicht_blinder_kampf_")):
            continue
        info_cell = ws.cell(row=r, column=5)
        if info_cell.value and "Zwölfstufige" in info_cell.value:
            info_cell.value = info_cell.value.replace("Zwölfstufige", "Achtzehnstufige")
            updated += 1
    assert updated == 12, f"Erwartete 12 aktualisierte Info-Texte, waren {updated}"

    # 2) 6 neue Stufen 13-18 anhaengen.
    assert ws.max_row == 2231, f"Erwartete letzte Zeile 2231, gefunden {ws.max_row}"
    for grad in range(13, 19):
        r = 2231 + (grad - 12)
        assert ws.cell(row=r, column=1).value is None, f"Zeile {r} ist nicht leer"
        vorher_grad = grad - 1
        ws.cell(row=r, column=1).value = f"vn_sicht_blinder_kampf_{grad}"
        ws.cell(row=r, column=2).value = "Vor- und Nachteile"
        ws.cell(row=r, column=3).value = f"Sicht: Blinder Kampf {ROEMISCH[grad]}"
        ws.cell(row=r, column=5).value = "Achtzehnstufige Vorteilskette; kein Blindheitsrabatt auf diese Stufe."
        ws.cell(row=r, column=6).value = "Blinder Kampf"
        ws.cell(row=r, column=7).value = "Auswahl"
        ws.cell(row=r, column=10).value = (
            f"VORAUSSETZUNG=vn_sicht_blinder_kampf_{vorher_grad} | BLINDER_KAMPF_STUFE={grad}"
        )
        ws.cell(row=r, column=11).value = grad
        ws.cell(row=r, column=12).value = 25
        ws.cell(row=r, column=13).value = "E"
        ws.cell(row=r, column=17).value = (
            f"Voraussetzung Blinder Kampf {ROEMISCH[vorher_grad]}. Durch Blindheit und "
            f"Lichtverhältnisse verursachte Sichterschwerungen werden um insgesamt {grad} "
            "reduziert, höchstens bis 0. Andere Sichterschwerungen gelten unverändert."
        )

    wb.save(XLSX_PATH)
    print("Blinder-Kampf-Kette auf 18 Stufen erweitert (Zeilen 2232-2237), 12 Info-Texte aktualisiert.")


if __name__ == "__main__":
    sys.exit(main())
