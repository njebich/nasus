"""LLM-007: die 7 leeren generischen Krankheit-Platzhalter aus dem Katalog entfernen.

DEC-1102 (VN-078, Owner-Durchgang 2026-09-10): vn_krankheit_i .. vii (Kosten 0, keine
Wirkung, keine Verfuegbarkeit) - "aktuell weder waehlbar noch Teil des Systems".
vn_krankheit_asthma (VN-013, Kurzatmigkeit) bleibt bestehen.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))

from openpyxl import load_workbook

WERTE = ROOT / "werte 0.8-claude.xlsx"
REFERENCES = [
    "vn_krankheit_i", "vn_krankheit_ii", "vn_krankheit_iii", "vn_krankheit_iv",
    "vn_krankheit_v", "vn_krankheit_vi", "vn_krankheit_vii",
]


def main():
    workbook = load_workbook(WERTE, read_only=False, data_only=False)
    worksheet = workbook["Werte"]
    headers = {cell.value: cell.column for cell in worksheet[1] if cell.value}

    rows = []
    for row in range(2, worksheet.max_row + 1):
        reference = worksheet.cell(row, headers["Referenz"]).value
        if reference in REFERENCES:
            rows.append((row, reference))

    if len(rows) != len(REFERENCES):
        raise RuntimeError(f"Erwartet {len(REFERENCES)} Zeilen, gefunden {len(rows)}: {rows}")

    row_numbers = [r for r, _ in rows]
    if row_numbers != list(range(row_numbers[0], row_numbers[0] + len(row_numbers))):
        raise RuntimeError(f"Zeilen nicht zusammenhaengend: {row_numbers}")

    worksheet.delete_rows(row_numbers[0], len(row_numbers))
    workbook.save(WERTE)
    print(f"Geloescht: Zeilen {row_numbers[0]}-{row_numbers[-1]} ({[ref for _, ref in rows]})")


if __name__ == "__main__":
    main()
