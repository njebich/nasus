# Material-Kanon-Korrektur (Nutzer 2026-09-12, "wir gehen jedes Material einmal durch, Material
# fuer Material"). Reine Volk-Datenkorrektur bestehender Zeilen (NK-Material + Schild-Material) -
# keine neuen Materialien/Ruestungstypen erfunden ("wir erfinden jetzt keine neuen Dinge"; ob eine
# passende Waffen-/Ruestungs-/Schild-Zeile fuer ein Material existiert, ist ausdruecklich NICHT
# Thema dieser Korrektur). Wird iterativ erweitert, waehrend der Nutzer Material fuer Material
# durchgeht - re-run nach jeder neuen Zeile in VOLK_FIXES.
import sys
import openpyxl

PATH = sys.argv[1] if len(sys.argv) > 1 else "werte 0.8-claude.xlsx"

# Materialname -> neuer Volk-Wert (Komma-Liste kanonischer Pluralnamen, oder "ALLE").
VOLK_FIXES = {
    "Bronze": "Katzen, Indianer",
    "Stein": "Indianer",
    "Eisen": "Dalkini, Draw, Elfen, Goblins, Orks, Trolle, Zentauren, Zwerge",  # alle außer Katzen/Indianer/Gnome
    "Qualitätsstahl": "Dalkini, Draw, Elfen, Goblins, Orks, Zwerge",  # wie Stahl, aber ohne Trolle
    "Qualitaetsstahl": "Dalkini, Draw, Elfen, Goblins, Orks, Zwerge",  # Schild-Material: andere Schreibweise (kein Umlaut)
    "Faltstahl": "Dalkini, Draw, Elfen, Goblins, Orks, Zwerge",  # ohne Trolle
    "Adamandit": "Draw, Elfen, Goblins, Orks, Zwerge",  # ohne Dalkini, Trolle (M-Sperre bleibt)
    "Chitin": "Katzen, Draw, Zwerge, Indianer",
    "Feineisen": "Dalkini, Draw, Elfen, Goblins, Orks, Trolle, Zentauren, Zwerge",  # Nutzer: "Feineisen = Eisen"
}

# Alte Materialnamen-Schreibweise -> korrekte Schreibweise (reine Umbenennung, keine Volk-Aenderung).
NAME_FIXES = {
    "Alchemistensilb.": "Alchemistensilber",
}


def find_col(ws, header):
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == header:
            return c
    raise KeyError(header)


def apply_sheet(wb, sheet_name, name_col, volk_col):
    ws = wb[sheet_name]
    changed = 0
    renamed = 0
    for r in range(2, ws.max_row + 1):
        raw_name = ws.cell(row=r, column=name_col).value
        if raw_name is None:
            continue
        name = raw_name.strip()
        if name in NAME_FIXES:
            ws.cell(row=r, column=name_col, value=NAME_FIXES[name])
            renamed += 1
            name = NAME_FIXES[name]
        if name in VOLK_FIXES:
            ws.cell(row=r, column=volk_col, value=VOLK_FIXES[name])
            changed += 1
    print(f"{sheet_name}: {changed} Volk-Korrektur(en), {renamed} Umbenennung(en)")


def main():
    wb = openpyxl.load_workbook(PATH)
    apply_sheet(wb, "NK-Material", find_col(wb["NK-Material"], "Waffe"), 11)
    apply_sheet(wb, "Schild-Material", find_col(wb["Schild-Material"], "Material"), find_col(wb["Schild-Material"], "Volk"))
    wb.save(PATH)


if __name__ == "__main__":
    main()
