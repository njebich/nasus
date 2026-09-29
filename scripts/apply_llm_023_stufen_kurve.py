"""LLM-023 (DEC-1295): Spruchmagie-Details Stufe1/2/3 auf die grad-basierte Kurve normalisieren.

Ausgenommen (nicht anfassen, eigene Owner-Entscheidungen): Wiederbelebung, Eisenhaut,
Korpi ignorieren, Durchblick. Erwartet genau 109 Aenderungen laut Log.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))
from openpyxl import load_workbook
from backup_werte import backup

WERTE = ROOT / "werte 0.8-claude.xlsx"
CURVE = {1: ("2", "6", "12"), 2: ("4", "9", "16"), 3: ("6", "12", "20"), 4: ("8", "15", "24"),
         5: ("9", "17", "27"), 6: ("10", "19", "30"), 7: ("11", "21", "33")}
EXCLUDED = {"spruchmagie_heilung_7_wiederbelebung", "spruchmagie_verzauberung_6_eisenhaut",
            "spruchmagie_antimagie_5_korpi_ignorieren", "spruchmagie_hellsicht_2_durchblick"}


def main():
    wb = load_workbook(WERTE)
    werte = wb["Werte"]
    wh = {c.value: c.column for c in werte[1] if c.value}
    grad = {werte.cell(r, wh["Referenz"]).value: int(werte.cell(r, wh["Grad"]).value)
            for r in range(2, werte.max_row + 1)
            if werte.cell(r, wh["Kategorie"]).value == "Spruchmagie" and werte.cell(r, wh["Referenz"]).value}
    ws = wb["Spruchmagie-Details"]
    h = {c.value: c.column for c in ws[1] if c.value}
    cols = [h["Stufe1"], h["Stufe2"], h["Stufe3"]]
    changes = []
    for r in range(2, ws.max_row + 1):
        ref = ws.cell(r, h["Referenz"]).value
        if not ref or ref in EXCLUDED:
            continue
        cur = tuple(str(ws.cell(r, c).value) for c in cols)
        tgt = CURVE[grad[ref]]
        if cur != tgt:
            changes.append((r, ref, cur, tgt))
    if len(changes) != 109:
        raise SystemExit(f"Erwartet 109 Aenderungen laut Log, gefunden {len(changes)}")
    print("Backup:", backup(str(WERTE)))
    for r, ref, cur, tgt in changes:
        for c, v in zip(cols, tgt):
            ws.cell(r, c, v)
        print(f"Z.{r} {ref}: {'/'.join(cur)} -> {'/'.join(tgt)}")
    wb.save(WERTE)


if __name__ == "__main__":
    main()
