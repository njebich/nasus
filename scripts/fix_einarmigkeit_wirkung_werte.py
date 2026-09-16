"""LLM-005-Nachtrag: Korrektur des wirkung-Texts von vn_koerper_einarmigkeit.

Die beim Anlegen (add_koerper_sprech_merkmale_werte.py) neu formulierte
Wirkung engte "einhaendige Waffenfuehrung" (RC-065 Sec 4.1, unqualifiziert)
faelschlich auf "einhaendige Nahkampfwaffen" ein. Owner-Korrektur
2026-09-13: gilt fuer einhaendige Waffen allgemein, nicht nur Nahkampf.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))

from openpyxl import load_workbook

WERTE = ROOT / "werte 0.8-claude.xlsx"
REFERENZ = "vn_koerper_einarmigkeit"
OLD = "Einhändige Nahkampfwaffen können ohne zusätzlichen Malus geführt werden."
NEW = "Einhändige Waffen können ohne zusätzlichen Malus geführt werden."


def main():
    workbook = load_workbook(WERTE, read_only=False, data_only=False)
    worksheet = workbook["Werte"]
    headers = {cell.value: cell.column for cell in worksheet[1] if cell.value}

    target_row = None
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row, headers["Referenz"]).value == REFERENZ:
            target_row = row
            break
    if target_row is None:
        raise RuntimeError(f"Referenz nicht gefunden: {REFERENZ}")

    wirkung_col = headers["Wirkung"]
    current = worksheet.cell(target_row, wirkung_col).value
    if OLD not in current:
        raise RuntimeError(f"Erwarteter Textbaustein nicht gefunden in Zeile {target_row}")

    updated = current.replace(OLD, NEW)
    worksheet.cell(target_row, wirkung_col, updated)
    workbook.save(WERTE)
    print(f"row={target_row} referenz={REFERENZ}")
    print(f"old={current}")
    print(f"new={updated}")


if __name__ == "__main__":
    main()
