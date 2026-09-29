"""LLM-030 (DEC-1392): alle 36 vn_meisterlicher_*_magier_{1,2,3}-Zeilen - Verfuegbarkeit
'3'/'5'/'7' (Kreis-Bindung) auf den NPC-only-Wert 'M' (nur Meister) setzen. Flag/Voraussetzungskette
bleiben unveraendert.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))
from openpyxl import load_workbook
from backup_werte import backup

WERTE = ROOT / "werte 0.8-claude.xlsx"


def main():
    wb = load_workbook(WERTE)
    ws = wb["Werte"]
    h = {c.value: c.column for c in ws[1] if c.value}
    col = h["Verfügbarkeit"] if "Verfügbarkeit" in h else h["Verfuegbarkeit"]
    targets = []
    for r in range(2, ws.max_row + 1):
        ref = ws.cell(r, h["Referenz"]).value or ""
        if ref.startswith("vn_meisterlicher_") and ref[-8:-2] == "magier" and ref[-1] in "123":
            targets.append((r, ref))
    if len(targets) != 36:
        raise SystemExit(f"Erwartet 36 Zeilen, gefunden {len(targets)}")
    expected = {"1": "3", "2": "5", "3": "7"}
    for r, ref in targets:
        if str(ws.cell(r, col).value) != expected[ref[-1]]:
            raise SystemExit(f"Ist-Stand weicht ab: {ref} = {ws.cell(r, col).value!r}")
    print("Backup:", backup(str(WERTE)))
    for r, ref in targets:
        ws.cell(r, col, "M")
    print(f"{len(targets)} Zeilen (Z.{targets[0][0]}..Z.{targets[-1][0]}) auf 'M' gesetzt")
    wb.save(WERTE)


if __name__ == "__main__":
    main()
