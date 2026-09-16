"""Fasst die neu hinzugekommene Warengruppe "Spielwaren" (3 Zeilen, Sheet "Preisliste") mit der
bereits bestehenden Warengruppe "Spielzeug" zusammen - inhaltliche Ueberschneidung
(Spielkarten/praeparierte Wuerfel passen zu den bereits vorhandenen Spielzeug-Wuerfeln),
Nutzer-Entscheidung 2026-09-16. Reine Kategorie-Umbenennung, keine sonstigen Feldaenderungen.
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

XLSX_PATH = Path(r"E:\Das Western Rollenspiel\LLM\werte 0.8-claude.xlsx")


def main():
    backup_path = XLSX_PATH.with_name(
        f"werte 0.8-claude_backup_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    )
    shutil.copy2(XLSX_PATH, backup_path)
    print(f"Backup angelegt: {backup_path}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    ws = wb["Preisliste"]

    renamed = 0
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=1)
        if cell.value == "Spielwaren":
            cell.value = "Spielzeug"
            renamed += 1
    assert renamed == 3, f"Erwartete 3 umbenannte Zeilen, waren {renamed}"

    wb.save(XLSX_PATH)
    print(f"{renamed} Zeilen von 'Spielwaren' auf 'Spielzeug' umbenannt.")


if __name__ == "__main__":
    sys.exit(main())
