"""LLM-014: Meisterliche-Grundfertigkeit-Katalogmigration.

DEC-330/VN-068: die 20 vn_faehigkeit_meisterliche_grundfertigkeit_*-Zeilen (Z. 1354-1373) zu
EINEM generischen Eintrag mit offener Zielauswahl zusammenfassen, analog zur bereits
vorhandenen dynamischen vn_faehigkeit_meisterliche_WHK-Zeile (Z. 1079, Flag
APP_DROPDOWN_WHK_OFFENE_LISTE) - reine Struktur-/Umsetzungsfrage, keine Regeländerung. Die
Grundfertigkeit-eigenen Zahlen (4 Mana pro TaW, (M)*2 min Dauer - abweichend von der WHK-Zeile
mit 2 Mana/TaW und (M) h) bleiben unveraendert erhalten.

Kein App-Code referenziert die 20 Einzelzeilen oder das WHK-Flag (grep-geprueft) - die "offene
Dropdown-Liste" ist bei der WHK-Zeile selbst noch nicht umgesetzt (Info-Text: "Umsetzung durch
Claude"), nur als Daten-Marker vorbereitet. Dieses Skript bringt die Grundfertigkeit-Zeile auf
denselben (noch unimplementierten) Stand, keinen weiteren.
"""
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

XLSX_PATH = Path(r"E:\Das Western Rollenspiel\LLM\werte 0.8-claude.xlsx")

FIRST_ROW = 1354
LAST_ROW = 1373


def main():
    backup_path = XLSX_PATH.with_name(
        f"werte 0.8-claude_backup_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    )
    shutil.copy2(XLSX_PATH, backup_path)
    print(f"Backup angelegt: {backup_path}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    ws = wb["Werte"]

    first_ref = ws.cell(row=FIRST_ROW, column=1).value
    last_ref = ws.cell(row=LAST_ROW, column=1).value
    assert first_ref == "vn_faehigkeit_meisterliche_grundfertigkeit_betoeren", first_ref
    assert last_ref == "vn_faehigkeit_meisterliche_grundfertigkeit_werfen", last_ref

    # Zeilen 1355-1373 (19 der 20) loeschen, Zeile 1354 wird unten zur neuen generischen Zeile.
    ws.delete_rows(FIRST_ROW + 1, LAST_ROW - FIRST_ROW)

    r = FIRST_ROW
    ws.cell(row=r, column=1).value = "vn_faehigkeit_meisterliche_grundfertigkeit"
    ws.cell(row=r, column=2).value = "Vor- und Nachteile"
    ws.cell(row=r, column=3).value = "Fähigkeit: Meisterliche Grundfertigkeit"
    ws.cell(row=r, column=5).value = (
        "Offene Grundfertigkeiten-Liste: In der App ist ein dynamisches Dropdown über alle "
        "vorhandenen Grundfertigkeiten erforderlich. Jede gewählte Grundfertigkeit ist eine "
        "eigene, separat kaufbare Instanz; dieselbe Grundfertigkeit darf nicht doppelt gewählt "
        "werden. Umsetzung durch Claude."
    )
    ws.cell(row=r, column=6).value = "Meisterliche Grundfertigkeit"
    ws.cell(row=r, column=7).value = "Auswahl"
    ws.cell(row=r, column=10).value = "APP_DROPDOWN_GRUNDFERTIGKEIT_OFFENE_LISTE"
    ws.cell(row=r, column=12).value = 100
    ws.cell(row=r, column=13).value = "E&3"
    ws.cell(row=r, column=17).value = (
        "Nicht Magier: Der Charakter kann max. (M) Mana für (M)*2 min in 1 Talentwert (TaW) "
        "pro 4 Mana einer ausgewählten Grundfertigkeit umwandeln. Dieser Vorteil muss für jede "
        "Grundfertigkeit einzeln gekauft werden."
    )

    wb.save(XLSX_PATH)
    print("20 vn_faehigkeit_meisterliche_grundfertigkeit_*-Zeilen zu 1 generischer Zeile zusammengefasst.")


if __name__ == "__main__":
    sys.exit(main())
