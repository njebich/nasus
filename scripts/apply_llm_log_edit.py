"""Setzt einzelne Eintraege aus nasus-nasus/docs/foundry-data-model/llm-anpassungen-log.md
(LLM-025..LLM-029) in der SPOT-xlsx um. spruchmagie.jsonl und spruchmagieDetails.json sind
Codegen-Artefakte von generate_data_ts.py - Aenderungen gehoeren in die xlsx.

Aufruf:
    python scripts/apply_llm_log_edit.py LLM-025

Legt vorher ein Backup nach backups/ an (scripts/backup_werte.py) und bricht ohne Aenderung ab,
wenn der Ist-Stand nicht dem im Log beschriebenen entspricht.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook

from backup_werte import backup

WERTE = ROOT / "werte 0.8-claude.xlsx"

GOLEM_SATZ = " Die Kontrollreichweite ist {M}*20m."

# (Sheet, Referenz, Spalte, alt, neu)
EDITS = {
    "LLM-025": [(
        "Werte", "spruchmagie_beherrschung_5_paralyse", "Wirkung",
        "Das Opfer erstarrt während der Wirkungsdauer vollständig ist aber bewegbar.",
        "Das Opfer erstarrt während der Wirkungsdauer vollständig (Status: Wehrlos), ist aber "
        "bewegbar. Steht das Opfer beim Eintritt der Wirkung, fällt es in der nächsten KR "
        "ungeprobt durch Status: Sturz.",
    )],
    "LLM-026": [(
        "Werte", "spruchmagie_beherrschung_3_tick_verursachen", "Wirkung",
        "Spezialisierbare Grund.- oder Optionalfertigkeit Aussuchen",
        "Spezialisierbare Grundfertigkeit oder WHK Aussuchen",
    )],
    "LLM-027": [
        ("Spruchmagie-Details", ref, "Gegenprobe", "alle 10 sec eine weitere", "alle att_magie sec eine weitere")
        for ref in (
            "spruchmagie_beherrschung_5_paralyse", "spruchmagie_beherrschung_6_totale_kontrolle",
            "spruchmagie_beherrschung_3_konfus", "spruchmagie_beherrschung_3_zorn",
        )
    ],
    "LLM-028": [(
        "Werte", "spruchmagie_beherrschung_6_totale_kontrolle", "Wirkung",
        "Wacht er auf und der Zauber fällt.",
        "Wacht er auf und der Zauber fällt.    Kontrollreichweite= {M} *10m",
    )],
    "LLM-029": [
        ("Werte", "spruchmagie_verzauberung_5_golem_kaempfer", "Wirkung",
         "Springen: max. 1m hoch, 6m weit.", "Springen: max. 1m hoch, 6m weit." + GOLEM_SATZ),
        ("Werte", "spruchmagie_verzauberung_7_grosser_golem", "Wirkung",
         "Der Golem gilt als Magiekorpus.", "Der Golem gilt als Magiekorpus." + GOLEM_SATZ),
    ],
}


def main():
    entry = sys.argv[1]
    edits = EDITS[entry]
    workbook = load_workbook(WERTE, read_only=False, data_only=False)

    plan = []
    for sheet, referenz, column, old, new in edits:
        ws = workbook[sheet]
        headers = {c.value: c.column for c in ws[1] if c.value}
        row = next((r for r in range(2, ws.max_row + 1)
                    if ws.cell(r, headers["Referenz"]).value == referenz), None)
        if row is None:
            raise SystemExit(f"{entry}: Referenz {referenz} nicht in Sheet {sheet}")
        cell = ws.cell(row, headers[column])
        current = cell.value or ""
        if old not in current:
            raise SystemExit(f"{entry}: Ist-Stand weicht ab bei {referenz} ({column}): {current!r}")
        if current.count(old) != 1:
            raise SystemExit(f"{entry}: Textbaustein nicht eindeutig bei {referenz}")
        if new in current:
            raise SystemExit(f"{entry}: {referenz} ist bereits umgesetzt")
        plan.append((cell, current.replace(old, new), sheet, referenz, row, current))

    dest = backup(str(WERTE))
    print(f"Backup: {dest}")
    for cell, updated, sheet, referenz, row, current in plan:
        cell.value = updated
        print(f"{sheet} Z.{row} {referenz}\n  alt: {current}\n  neu: {updated}")
    workbook.save(WERTE)


if __name__ == "__main__":
    main()
