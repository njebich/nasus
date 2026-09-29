"""LLM-024 (DEC-1321): NK-Waffen-Basis, Hauptfertigkeit/Spezialisierung Unbewaffnet:
Eisenkappenstiefel, Orkischer Kampfhandschuh, Panzerschienen, Totschlaeger auskommentieren
('#'-Praefix am Waffe-Namen), 'Panzerhandschuhe' -> 'Gepanzert' umbenennen.
Die gleichnamige Hiebwaffen-Zeile 'Totschlaeger' (Stumpfe Hiebwaffen) bleibt unveraendert.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))
from openpyxl import load_workbook
from backup_werte import backup

WERTE = ROOT / "werte 0.8-claude.xlsx"
COMMENT_OUT = ["Eisenkappenstiefel", "Orkischer Kampfhandschuh", "Panzerschienen", "Totschläger"]


def main():
    wb = load_workbook(WERTE)
    ws = wb["NK-Waffen-Basis"]
    h = {c.value: c.column for c in ws[1] if c.value}
    rows = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, h["Hauptfertigkeit"]).value == "Unbewaffnet" and ws.cell(r, h["Spezialisierung"]).value == "Unbewaffnet":
            rows.setdefault(str(ws.cell(r, h["Waffe"]).value).strip(), []).append(r)
    for name in COMMENT_OUT + ["Panzerhandschuhe"]:
        if len(rows.get(name, [])) != 1:
            raise SystemExit(f"Ist-Stand weicht ab: {name} -> {rows.get(name)}")
    print("Backup:", backup(str(WERTE)))
    for name in COMMENT_OUT:
        r = rows[name][0]
        ws.cell(r, h["Waffe"], "#" + ws.cell(r, h["Waffe"]).value.strip())
        print(f"Z.{r} auskommentiert: {name}")
    r = rows["Panzerhandschuhe"][0]
    ws.cell(r, h["Waffe"], "Gepanzert")
    print(f"Z.{r} umbenannt: Panzerhandschuhe -> Gepanzert")
    wb.save(WERTE)


if __name__ == "__main__":
    main()
