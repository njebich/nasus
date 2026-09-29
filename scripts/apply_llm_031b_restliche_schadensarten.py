"""LLM-031 Nachtrag (Owner-Entscheidung 2026-09-29): Armklingen-Familie (Klauen) = Scharf,
exotische Kampfstile = Stumpf. Fuellt die 11 zuvor leeren Schadensart-Zellen in NK-Waffen-Basis.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))
from openpyxl import load_workbook
from backup_werte import backup

WERTE = ROOT / "werte 0.8-claude.xlsx"
KLAUEN = {"Drei Langklauen", "Klaue", "Pranke", "Zwei Klauen"}
KAMPFSTILE = {"Boxen", "Elfische Kunst der Selbstverteidigung", "Goblinische Kampfkunst",
              "Katzenmenschen Kampfkunst", "Orkisch' Raufen", "Ringen", "Schattenkampf"}


def main():
    wb = load_workbook(WERTE)
    ws = wb["NK-Waffen-Basis"]
    h = {c.value: c.column for c in ws[1] if c.value}
    todo = []
    for r in range(2, ws.max_row + 1):
        name, spez = ws.cell(r, h["Waffe"]).value, ws.cell(r, h["Spezialisierung"]).value
        if ws.cell(r, h["Schadensart"]).value:
            continue
        if spez == "Armklingen" and name in KLAUEN:
            todo.append((r, name, "Scharf"))
        elif spez in KAMPFSTILE and name == spez:
            todo.append((r, name, "Stumpf"))
    if len(todo) != 11:
        raise SystemExit(f"Erwartet 11 leere Zeilen, gefunden {len(todo)}")
    print("Backup:", backup(str(WERTE)))
    for r, name, val in todo:
        ws.cell(r, h["Schadensart"], val)
        print(f"Z.{r} {name}: {val}")
    wb.save(WERTE)


if __name__ == "__main__":
    main()
