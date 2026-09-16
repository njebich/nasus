"""LLM-005: die fuenf Koerper-/Sprech-Merkmale aus DEC-137 (RC-065, DEC-1101) als neue
V/N-Katalogzeilen anlegen. Bisher kein Datensatz vorhanden.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))

from openpyxl import load_workbook

WERTE = ROOT / "werte 0.8-claude.xlsx"

ENTRIES = [
    {
        "Referenz": "vn_koerper_einarmigkeit",
        "Kategorie": "Vor- und Nachteile",
        "Beschreibung": "Körper: Einarmigkeit",
        "Art": "Auswahl",
        "Kosten": -400,
        "Verfuegbarkeit": "E",
        "Wirkung": (
            "Dem Charakter fehlt von Geburt an ein Arm. Er erleidet dadurch eine permanente "
            "Behinderung von 3 (KBE) auf körperliche Proben; auf geistige Proben (MBE) hat der "
            "fehlende Arm keinen Einfluss. Einhändige Nahkampfwaffen können ohne zusätzlichen "
            "Malus geführt werden. Der verbleibende Arm genügt für alle nötigen Gesten; die "
            "Spruchmagie ist dadurch nicht eingeschränkt."
        ),
    },
    {
        "Referenz": "vn_koerper_einbeinigkeit",
        "Kategorie": "Vor- und Nachteile",
        "Beschreibung": "Körper: Einbeinigkeit",
        "Art": "Auswahl",
        "Kosten": -400,
        "Verfuegbarkeit": "E",
        "Wirkung": (
            "Dem Charakter fehlt von Geburt an ein Bein. Er erleidet dadurch eine permanente "
            "Behinderung von 4 (KBE) auf körperliche Proben; auf geistige Proben (MBE) hat das "
            "fehlende Bein keinen Einfluss. Ohne Hilfsmittel oder Hilfe bewegt er sich "
            "ausschließlich kriechend mit fester Geschwindigkeit von 1 Feld pro Kampfrunde; "
            "Joggen und Sprint sind nicht möglich. Mit Hilfsmitteln stehen alle "
            "Fortbewegungsarten offen: eine einfache Krücke oder ein Holzbein multipliziert jede "
            "Geschwindigkeit mit 0,5, eine angepasste Gehhilfe oder tatkräftige fremde Hilfe mit "
            "0,75 (danach aufgerundet)."
        ),
    },
    {
        "Referenz": "vn_koerper_keine_arme",
        "Kategorie": "Vor- und Nachteile",
        "Beschreibung": "Körper: Keine Arme",
        "Art": "Auswahl",
        "Kosten": -800,
        "Verfuegbarkeit": "E",
        "Wirkung": (
            "Dem Charakter fehlen von Geburt an beide Arme. Er erleidet dadurch eine permanente "
            "Behinderung von 6 (KBE) auf körperliche Proben; auf geistige Proben (MBE) haben die "
            "fehlenden Arme keinen Einfluss. Er kann keine Nah-, Fern- oder Wurfwaffe und keinen "
            "Schild führen. Spruchmagie kann er nur auf Zauberstufe 2 oder 3 wirken (keine Geste "
            "möglich): auf Stufe 2 entfällt die Geste, die Formel bleibt erforderlich; Stufe 1 "
            "ist ausgeschlossen. Ohne das jeweilige Talent (talente_spruchmagie_stufe_2_zaubern "
            "bzw. _stufe_3_zaubern) kann er faktisch nicht zaubern. In Kombination mit Stumm ist "
            "nur noch Zauberstufe 3 möglich (weder Geste noch Formel; "
            "talente_spruchmagie_stufe_3_zaubern erforderlich). Fortbewegung ist ansonsten "
            "uneingeschränkt; Kriechen setzt jedoch mindestens zwei funktionsfähige Gliedmaßen "
            "voraus."
        ),
    },
    {
        "Referenz": "vn_koerper_keine_beine",
        "Kategorie": "Vor- und Nachteile",
        "Beschreibung": "Körper: Keine Beine",
        "Art": "Auswahl",
        "Kosten": -800,
        "Verfuegbarkeit": "E",
        "Wirkung": (
            "Dem Charakter fehlen von Geburt an beide Beine. Er erleidet dadurch eine permanente "
            "Behinderung von 8 (KBE) auf körperliche Proben; auf geistige Proben (MBE) haben die "
            "fehlenden Beine keinen Einfluss. Er bewegt sich ausschließlich kriechend mit fester "
            "Geschwindigkeit von 1 Feld pro Kampfrunde, sofern mindestens zwei funktionsfähige "
            "Gliedmaßen vorhanden sind; Joggen und Sprint sind nicht möglich. Hilfsmittel können "
            "diese Geschwindigkeit wie bei Einbeinigkeit abstufen, soweit anwendbar."
        ),
    },
    {
        "Referenz": "vn_sprech_stumm",
        "Kategorie": "Vor- und Nachteile",
        "Beschreibung": "Sprech: Stumm",
        "Art": "Auswahl",
        "Kosten": -300,
        "Verfuegbarkeit": "E",
        "Wirkung": (
            "Der Charakter kann von Geburt an keinen Laut hervorbringen (nur die Atemluft bleibt "
            "hörbar) — identisch mit der Wirkung des Zaubers Stumm, jedoch dauerhaft. Er trägt "
            "dadurch keine Behinderung (KBE 0, MBE 0). Spruchmagie kann er nur ab Zauberstufe 2 "
            "wirken (keine Formel möglich, Geste bleibt erforderlich); ohne das jeweilige Talent "
            "(talente_spruchmagie_stufe_2_zaubern bzw. _stufe_3_zaubern) kann er faktisch nicht "
            "zaubern. In Kombination mit Keine Arme ist nur noch Zauberstufe 3 möglich (weder "
            "Geste noch Formel; talente_spruchmagie_stufe_3_zaubern erforderlich). Fortbewegung "
            "ist uneingeschränkt."
        ),
    },
]


def main():
    workbook = load_workbook(WERTE, read_only=False, data_only=False)
    worksheet = workbook["Werte"]
    headers = {cell.value: cell.column for cell in worksheet[1] if cell.value}

    existing = {
        worksheet.cell(row, headers["Referenz"]).value
        for row in range(2, worksheet.max_row + 1)
    }
    for entry in ENTRIES:
        if entry["Referenz"] in existing:
            raise RuntimeError(f"Referenz existiert bereits: {entry['Referenz']}")

    written = []
    for entry in ENTRIES:
        row = worksheet.max_row + 1
        for header, value in entry.items():
            worksheet.cell(row, headers[header], value)
        written.append((row, entry["Referenz"]))

    workbook.save(WERTE)
    print(f"written={len(written)}")
    for row, reference in written:
        print(row, reference)


if __name__ == "__main__":
    main()
