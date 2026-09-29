"""LLM-020: Feuerwaffen-Sheet - 23 Zeilen auf den Typ-Standard-NK-Improvisiert-Block umstellen.

Gewehr -> Stangenwaffen/"Improvisierte Stangenwaffen" (11 Zeilen waren Hiebwaffen),
Pistole -> Hiebwaffen/"Improvisierte Hiebwaffen" (12 Zeilen waren Stangenwaffen).
Blockwerte werden aus den bereits korrekten Zeilen des jeweiligen Typs abgeleitet;
Min-Staerke-1H/2H-Basis folgen der jeweiligen Min.-Staerke der Waffe (Gewehr 1H=2H, Pistole 2H='-').
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))
from openpyxl import load_workbook
from backup_werte import backup

WERTE = ROOT / "werte 0.8-claude.xlsx"
BLOCK_COLS = ["Hauptfertigkeit", "Spezialisierung", "WK-Basis", "Schadenswuerfel-1", "Schadenswuerfel-2",
              "Staerke-Teiler", "Staerke-Malus-Basis", "AT-Basis", "PA-Basis",
              "Klingenbrecher-Basis", "Klingenschutz-Basis"]
WRONG = {"Gewehr": "Hiebwaffen", "Pistole": "Stangenwaffen"}
RIGHT = {"Gewehr": "Stangenwaffen", "Pistole": "Hiebwaffen"}


def main():
    wb = load_workbook(WERTE)
    ws = wb["Feuerwaffen"]
    h = {c.value: c.column for c in ws[1] if c.value}
    minst = next(col for name, col in h.items() if name.startswith("Min. St"))
    rows = range(2, ws.max_row + 1)
    val = lambda r, name: ws.cell(r, h[name]).value

    template = {}
    for r in rows:
        typ = val(r, "Typ")
        if typ in RIGHT and val(r, "Hauptfertigkeit") == RIGHT[typ] and typ not in template:
            template[typ] = {c: val(r, c) for c in BLOCK_COLS}
    targets = [r for r in rows if val(r, "Typ") in WRONG and val(r, "Hauptfertigkeit") == WRONG[val(r, "Typ")]]
    counts = {t: sum(1 for r in targets if val(r, "Typ") == t) for t in WRONG}
    if counts != {"Gewehr": 11, "Pistole": 12}:
        raise SystemExit(f"Ist-Stand weicht vom Log ab: {counts}")

    print("Backup:", backup(str(WERTE)))
    for r in targets:
        typ = val(r, "Typ")
        for c in BLOCK_COLS:
            ws.cell(r, h[c], template[typ][c])
        m = ws.cell(r, minst).value
        ws.cell(r, h["Min-Staerke-1H-Basis"], m)
        ws.cell(r, h["Min-Staerke-2H-Basis"], m if typ == "Gewehr" else "-")
        print(f"Z.{r} {val(r, 'Name')} ({typ}) -> {RIGHT[typ]}")
    wb.save(WERTE)


if __name__ == "__main__":
    main()
