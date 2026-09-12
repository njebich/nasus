# One-off: Ruestung/Schild Voelkerzuweisung + Schild-Verfuegbarkeit (AW/NW), per den bestaetigten
# Grundentscheidungen 28/30/31/32/33 in nasus-spec/nasus-spec/04-Verfuegbarkeiten-und-
# Herkunftsorte.md. Ruestung-Basis hat bereits Verfuegbarkeit-AW/NW (frueherer Stand) - hier kommt
# nur die fehlende Volk-Spalte dazu. Schild-Material/-Fertigung/-Bespannung hatten bisher weder
# Verfuegbarkeit noch Volk - beides wird hier ergaenzt. Schild-Verplatung bleibt unberuehrt (reines
# Stub-Sheet ohne Werte, siehe shieldComposition.ts-Dateikopf).
# Schreibt NUR neue Spalten am Sheet-Ende an (keine Einfuegung) - bestehende formelbasierte
# Spaltenverweise anderswo im Workbook bleiben unberuehrt.
# Nutzer-bestaetigt (2026-09-12): kein Volk angegeben = alle Voelker - ALLE-Zeilen bleiben deshalb
# leer statt mit dem Literal "ALLE" befuellt (konsistent mit Spec-Punkt 20 und der bestehenden
# istWaffenKomponenteVerfuegbar-Pruefung: `!volk || volk === 'ALLE'`).
import sys
import openpyxl

PATH = sys.argv[1] if len(sys.argv) > 1 else "werte 0.8-claude.xlsx"

METALLLISTE_7 = "Dalkini, Draw, Elfen, Goblins, Orks, Trolle, Zwerge"


def col_letter(idx):
    return openpyxl.utils.get_column_letter(idx)


def append_columns(ws, names):
    """Haengt neue Header ans Sheet-Ende an, gibt {name: column_index} zurueck."""
    start = ws.max_column + 1
    out = {}
    for i, name in enumerate(names):
        c = start + i
        ws.cell(row=1, column=c, value=name)
        out[name] = c
    return out


def find_col(ws, header):
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == header:
            return c
    raise KeyError(header)


# ---------------------------------------------------------------------------
# Ruestung-Basis (Punkt 28): 6 Kettenruestungs- + 6 Metallplattenruestungsbasen bekommen eine
# AUSWAHL-Voelkerliste, die uebrigen 8 Basen sowie alle Verarbeitungs-/Anpassungsstufen bleiben
# ALLE (Verarbeitung/Anpassung brauchen deshalb gar keine Volk-Spalte).
KETTEN_RUESTUNGEN = {
    "Eisen-Kettenpanzer", "Stahl-Kettenpanzer", "Schwerer Eisen-Kettenpanzer",
    "Faltstahl-Kettenpanzer", "Schwerer Stahl-Kettenpanzer", "Schwerer Faltstahl-Kettenpanzer",
}
KETTEN_VOLK = "Dalkini, Draw, Elfen, Goblins, Orks, Zwerge"  # ohne Trolle (Punkt 28)

METALLPLATTEN_RUESTUNGEN = {
    "Eisenpanzer", "Stahlpanzer", "Schwerer Eisenpanzer", "Schwerer Stahlpanzer",
    "Faltstahlpanzer", "Schwerer Faltstahlpanzer",
}
METALLPLATTEN_VOLK = METALLLISTE_7

# ---------------------------------------------------------------------------
# Schild-Material (Punkt 30/32): globale Materialreferenz, Werte identisch zur bereits fuer
# NK-Material verwendeten Referenztabelle (add_nk_waffen_verfuegbarkeit.py), erweitert um die 3
# dort fehlenden Eintraege Feineisen/Spinnenwebe/Drachenschuppe. Schluessel exakt wie im Sheet
# geschrieben (nach den Punkt-30-Umbenennungen Kolhartz->Kolharz/Diamantspart->Diamantspat/
# Schwarzfells->Schwarzfels; "Hart Holz"/"Drachensch." werden NICHT umbenannt - Punkt 30 nennt nur
# die drei genannten Korrekturen, keine Freilizenz fuer weitere Namensangleichungen).
MATERIAL_RENAME = {"Kolhartz": "Kolharz", "Diamantspart": "Diamantspat", "Schwarzfells": "Schwarzfels"}
MATERIAL_AW_NW = {
    "Holz": (1, 1), "Hart Holz": (2, 2), "Leder": (1, 1), "Spinnenwebe": (4, 4), "Chitin": (5, 2),
    "Drachensch.": ("M", "M"), "Stein": (1, 1), "Schwarzfels": (5, 1), "Diamantspat": (7, 7),
    "Bronze": (1, 1), "Eisen": (1, 1), "Feineisen": (1, 1), "Stahl": (1, 2), "Qualitaetsstahl": (2, 4),
    "Faltstahl": (3, 5), "Mithril": (7, 7), "Adamandit": ("M", "M"), "Alchemistensilb.": (5, 6),
    "Vulkanglas": (7, 3), "Nasium": (7, 7), "Knochen": (1, 1), "Kolharz": (5, 7),
}
# Nur nicht-ALLE Eintraege noetig (fehlender Key = ALLE, siehe Dateikopf). Mithril/Nasium enger als
# die allgemeine Metallliste (nur Elfen+Zwerge), Alchemistensilber ohne Draw/Trolle (Punkt 32).
# Spinnenwebe->Draw hier bewusst wie bei der Schild-Bespannung (Punkt 33) behandelt, obwohl Punkt
# 30s eigene "18 Komponenten bleiben ALLE"-Zaehlung nur aufgeht, wenn man Spinnenwebe (Material)
# stattdessen als ALLE zaehlt - Punkt 32 ("Spinnenwebe ausschliesslich Draw", "gilt verbindlich fuer
# Waffen, Schilde") und Punkt 33 (Bespannung-Spinnenwebe: Draw) sind aber beide explizit und
# eindeutig; die Diskrepanz wird hier dokumentiert statt stillschweigend zugunsten der Zaehlung
# aufgeloest.
MATERIAL_VOLK = {
    "Mithril": "Elfen, Zwerge", "Nasium": "Elfen, Zwerge",
    "Alchemistensilb.": "Dalkini, Elfen, Goblins, Orks, Zwerge",
    "Diamantspat": "Draw", "Schwarzfels": "Indianer, Katzen", "Vulkanglas": "Indianer, Katzen",
    "Kolharz": "Zentauren", "Spinnenwebe": "Draw",
    "Eisen": METALLLISTE_7, "Feineisen": METALLLISTE_7, "Stahl": METALLLISTE_7,
    "Qualitaetsstahl": METALLLISTE_7, "Faltstahl": METALLLISTE_7, "Bronze": METALLLISTE_7,
    "Adamandit": METALLLISTE_7,
}

# ---------------------------------------------------------------------------
# Schild-Fertigung (Punkt 31): identisch zur bestaetigten NK-Fertigung. Nur "Goblin Massenfab."
# ist auf Goblins beschraenkt (Punkt 30 letzter Satz), alle anderen Fertigungen gelten fuer ALLE.
FERTIGUNG_AW_NW = {
    "Ausschuss": (1, 1), "Goblin Massenfab.": (1, 1), "Massenfabrikation": (1, 1),
    "Gesellenarbeit": (1, 1), "Meisterarbeit": (2, 3), "Großmeisterarbeit": (3, 5),
    "Einzelstück": (7, 7),
}
FERTIGUNG_VOLK = {"Goblin Massenfab.": "Goblins"}

# ---------------------------------------------------------------------------
# Schild-Bespannung (Punkt 33, vollstaendig bestaetigt).
BESPANNUNG_AW_NW = {
    "Stoff": (1, 1), "Leder": (1, 1), "Spinnenwebe": (4, 4), "Kohlharz": (4, 4),
    "Drachenschuppe": ("M", "M"),
}
BESPANNUNG_VOLK = {"Spinnenwebe": "Draw", "Kohlharz": "Zentauren"}


def main():
    wb = openpyxl.load_workbook(PATH)  # nicht read_only - wir schreiben.
    report = []

    # --- Ruestung-Basis ----------------------------------------------------
    ws = wb["Ruestung-Basis"]
    name_col = find_col(ws, "Ruestungsteil")
    cols = append_columns(ws, ["Volk"])
    kette = platte = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        if name is None:
            continue
        if name in KETTEN_RUESTUNGEN:
            ws.cell(row=r, column=cols["Volk"], value=KETTEN_VOLK)
            kette += 1
        elif name in METALLPLATTEN_RUESTUNGEN:
            ws.cell(row=r, column=cols["Volk"], value=METALLPLATTEN_VOLK)
            platte += 1
    report.append(f"Ruestung-Basis: {kette} Kettenruestungen, {platte} Metallplattenruestungen mit Volk-AUSWAHL")

    # --- Schild-Material -----------------------------------------------
    ws = wb["Schild-Material"]
    name_col = find_col(ws, "Material")
    cols = append_columns(ws, ["Verfuegbarkeit-AW", "Verfuegbarkeit-NW", "Volk"])
    volk_gesetzt = 0
    for r in range(2, ws.max_row + 1):
        raw_name = ws.cell(row=r, column=name_col).value
        if raw_name is None:
            continue
        stripped = raw_name.strip()
        name = MATERIAL_RENAME.get(stripped, stripped)
        if name != raw_name:
            ws.cell(row=r, column=name_col, value=name)
        aw, nw = MATERIAL_AW_NW[name]
        ws.cell(row=r, column=cols["Verfuegbarkeit-AW"], value=aw)
        ws.cell(row=r, column=cols["Verfuegbarkeit-NW"], value=nw)
        if name in MATERIAL_VOLK:
            ws.cell(row=r, column=cols["Volk"], value=MATERIAL_VOLK[name])
            volk_gesetzt += 1
    report.append(f"Schild-Material: {ws.max_row - 1} Zeilen mit Verfuegbarkeit, {volk_gesetzt} mit Volk-AUSWAHL")

    # --- Schild-Fertigung ------------------------------------------------
    ws = wb["Schild-Fertigung"]
    name_col = find_col(ws, "Fertigung")
    cols = append_columns(ws, ["Verfuegbarkeit-AW", "Verfuegbarkeit-NW", "Volk"])
    volk_gesetzt = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        if name is None:
            continue
        aw, nw = FERTIGUNG_AW_NW[name.strip()]
        ws.cell(row=r, column=cols["Verfuegbarkeit-AW"], value=aw)
        ws.cell(row=r, column=cols["Verfuegbarkeit-NW"], value=nw)
        if name.strip() in FERTIGUNG_VOLK:
            ws.cell(row=r, column=cols["Volk"], value=FERTIGUNG_VOLK[name.strip()])
            volk_gesetzt += 1
    report.append(f"Schild-Fertigung: {ws.max_row - 1} Zeilen mit Verfuegbarkeit, {volk_gesetzt} mit Volk-AUSWAHL")

    # --- Schild-Bespannung -------------------------------------------------
    ws = wb["Schild-Bespannung"]
    name_col = find_col(ws, "Bespannung")
    cols = append_columns(ws, ["Verfuegbarkeit-AW", "Verfuegbarkeit-NW", "Volk"])
    volk_gesetzt = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=name_col).value
        if name is None:
            continue
        aw, nw = BESPANNUNG_AW_NW[name.strip()]
        ws.cell(row=r, column=cols["Verfuegbarkeit-AW"], value=aw)
        ws.cell(row=r, column=cols["Verfuegbarkeit-NW"], value=nw)
        if name.strip() in BESPANNUNG_VOLK:
            ws.cell(row=r, column=cols["Volk"], value=BESPANNUNG_VOLK[name.strip()])
            volk_gesetzt += 1
    report.append(f"Schild-Bespannung: {ws.max_row - 1} Zeilen mit Verfuegbarkeit, {volk_gesetzt} mit Volk-AUSWAHL")

    wb.save(PATH)
    print("\n".join(report))


if __name__ == "__main__":
    main()
