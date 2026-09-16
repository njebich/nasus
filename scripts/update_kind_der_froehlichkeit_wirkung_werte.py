"""LLM-009: den Wirkungstext von vn_kind_der_froehlichkeit an RC-069 (DEC-1104) angleichen.

(1) KI-Klammer entfernen, KI-Passus zu "Maximum um 1/2/3/4 erhoeht" (KI-Maximum 24, keine
    Sonderregel mehr).
(2) "Optios und Sprachen" -> "WHK-Fertigkeiten"; SSK ersatzlos gestrichen.
(3) Wunder ergaenzt (genutztes Gebets-WHK-Talent).
(4) "Kampftalente" -> "Kampf-Hauptfertigkeiten" + eigene Ausweichen-Klausel.
(5) Ausloeser/Kaskade/Fehlschlag-Klausel ergaenzt (jede gewuerfelte 29).
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))

from openpyxl import load_workbook

WERTE = ROOT / "werte 0.8-claude.xlsx"
REFERENCE = "vn_kind_der_froehlichkeit"

WIRKUNG = (
    "Wirft der Charakter einen fröhlichen Erfolg (eine gewürfelte 29, auch bei reduziertem "
    "Probenwert) auf einer Probe mit Werttyp, hat er sofort Anspruch auf eine weitere Probe auf "
    "denselben Wert unter exakt denselben Bedingungen (reiner Wachstums-Wurf, keine eigene "
    "Aktion, kein Mana- oder Handlungsverbrauch). Bei normalem/gutem/meisterlichem/fröhlichem "
    "Erfolg in dieser Zusatzprobe erhöht sich das Maximum bei Eigenschaften um 0/1/2/3, bei "
    "Attributen um 0/0/0/1, bei Grundfertigkeiten um 1/2/3/4, bei Kampf-Hauptfertigkeiten "
    "(Nah-/Fernkampf) um 0/1/2/3, bei einer Ausweichen-Probe je um 0/1/2/3 auf sf_ausweichen, "
    "sf_gutes_aw und sf_meisterliches_ausweichen, bei WHK-Fertigkeiten um 2/4/6/8, bei Wundern "
    "wird das genutzte Gebets-WHK-Talent (whk_geweihte_stossgebet/_wunder/_ritual) um 2/4/6/8 "
    "erhöht (wie WHK), bei PSI- und Spruchzaubern (jeweils für den einzelnen Zauber) um 1/2/3/4; "
    "bei KI-Fähigkeiten wird das Maximum um 1/2/3/4 erhöht. Fällt in der Zusatzprobe eine "
    "erneute 29, kaskadiert der Effekt (weitere Zusatzprobe, Erhöhungen summieren sich). "
    "Fehlschlag oder Patzer der Zusatzprobe bedeuten keine Erhöhung und lösen keine "
    "Patzer-Folge aus. Jeder Spielercharakter besitzt diesen Vorteil automatisch."
)


def main():
    workbook = load_workbook(WERTE, read_only=False, data_only=False)
    worksheet = workbook["Werte"]
    headers = {cell.value: cell.column for cell in worksheet[1] if cell.value}
    row = next(
        (
            row
            for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row, headers["Referenz"]).value == REFERENCE
        ),
        None,
    )
    if row is None:
        raise RuntimeError(f"Referenz nicht gefunden: {REFERENCE}")

    worksheet.cell(row, headers["Wirkung"], WIRKUNG)
    workbook.save(WERTE)
    print(row, REFERENCE)


if __name__ == "__main__":
    main()
