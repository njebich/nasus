# Material-Kanon-Nachzug fuer NK-Schaftmaterial (2026-09-12), Fortsetzung von
# scripts/fix_material_kanon.py: die Schaftmaterial-Zeilen sind materialbenannt ("Eisen
# Verstaerkt"/"Eisen Voll" usw.), wurden beim urspruenglichen Material-Kanon-Pass aber
# nicht angefasst und tragen noch die alte pauschale Voelkerliste. Reine Volk-Korrektur
# anhand des bereits mit dem Nutzer abgestimmten Kanons (NK-Material/Schild-Material) -
# keine neuen Materialien/Zeilen erfunden.
import sys
import openpyxl

PATH = sys.argv[1] if len(sys.argv) > 1 else "werte 0.8-claude.xlsx"

# Materialname (Praefix des Zeilennamens vor "Verstaerkt"/"Voll") -> kanonischer Volk-Wert.
VOLK_FIXES = {
    "Eisen": "Dalkini, Draw, Elfen, Goblins, Orks, Trolle, Zentauren, Zwerge",
    "Stahl": "Dalkini, Draw, Elfen, Goblins, Orks, Trolle, Zwerge",  # unveraendert
    "Qualitätsstahl": "Dalkini, Draw, Elfen, Goblins, Orks, Zwerge",  # ohne Trolle
    "Faltstahl": "Dalkini, Draw, Elfen, Goblins, Orks, Zwerge",  # ohne Trolle
    "Mithril": "Elfen, Zwerge",  # unveraendert
    "Adamandit": "Draw, Elfen, Goblins, Orks, Zwerge",  # ohne Dalkini, Trolle
    "Bronze": "Katzen, Indianer",
}


def material_praefix(name):
    for material in VOLK_FIXES:
        if name.startswith(material):
            return material
    return None


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb["NK-Schaftmaterial"]
    name_col, volk_col = 1, 11
    changed = 0
    for r in range(2, ws.max_row + 1):
        raw_name = ws.cell(row=r, column=name_col).value
        if raw_name is None:
            continue
        material = material_praefix(raw_name.strip())
        if material is None:
            continue
        neu = VOLK_FIXES[material]
        if ws.cell(row=r, column=volk_col).value != neu:
            ws.cell(row=r, column=volk_col, value=neu)
            changed += 1
            print(f"  Zeile {r} ({raw_name.strip()!r}): -> {neu}")
    print(f"NK-Schaftmaterial: {changed} Volk-Korrektur(en)")
    wb.save(PATH)


if __name__ == "__main__":
    main()
