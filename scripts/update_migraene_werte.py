"""LLM-008: die 5 Migraene-Stufen an RC-068 (DEC-1103) angleichen.

(1) vn_krankheit_extrem_migraene: KON-Erschwerung 30 -> 25.
(2) Ruheintervall W3 Tage -> W8 Spieltage (alle 5 Stufen).
(3) Deckel 9 auf die akkumulierte Behinderung ergaenzen.
(4) EXKLUSIV_MIGRAENE-Flag ergaenzen (maximal eine Stufe pro Charakter).
(5) Wirkungstext "W8 BE" -> "W8 als KBE und MBE (Mentale Behinderung)",
    gedeckelt bei 9; die akkumulierte Behinderung erschwert die taegliche
    KON-Konterprobe nicht.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))

from openpyxl import load_workbook

WERTE = ROOT / "werte 0.8-claude.xlsx"

GRADE = [
    ("vn_krankheit_leicht_migraene", 5),
    ("vn_krankheit_mittel_migraene", 10),
    ("vn_krankheit_schwer_migraene", 15),
    ("vn_krankheit_sehr_schwer_migraene", 20),
    ("vn_krankheit_extrem_migraene", 25),
]

WIRKUNG_TEMPLATE = (
    "Nach W8 migränefreien Spieltagen legt der Charakter eine um {erschwerung} erschwerte "
    "KON-Probe ab. Gelingt sie, bleibt der Tag ohne Migräne, und am Folgetag beginnt eine neue "
    "Ruhephase von W8 migränefreien Spieltagen. Misslingt sie, erleidet der Charakter für diesen "
    "Tag einen ausgewürfelten W8 als KBE und MBE (Mentale Behinderung). An jedem unmittelbar "
    "folgenden Spieltag wird die weiterhin nur um {erschwerung} erschwerte KON-Probe erneut "
    "abgelegt; die bereits angesammelte Behinderung erschwert diese Konterprobe nicht. Bei jedem "
    "weiteren Fehlschlag wird ein neuer W8 gewürfelt und additiv zur bestehenden Behinderung "
    "addiert, gedeckelt bei 9. Der erste migränefreie Tag beendet die Anfallsphase sofort, setzt "
    "die Behinderung auf 0 zurück und startet am Folgetag eine neue Ruhephase."
)

FLAG = "EXKLUSIV_MIGRAENE=migraene"


def main():
    workbook = load_workbook(WERTE, read_only=False, data_only=False)
    worksheet = workbook["Werte"]
    headers = {cell.value: cell.column for cell in worksheet[1] if cell.value}
    rows_by_reference = {
        worksheet.cell(row, headers["Referenz"]).value: row
        for row in range(2, worksheet.max_row + 1)
    }

    updated = []
    for reference, erschwerung in GRADE:
        row = rows_by_reference.get(reference)
        if row is None:
            raise RuntimeError(f"Referenz nicht gefunden: {reference}")
        worksheet.cell(row, headers["Flag"], FLAG)
        worksheet.cell(row, headers["Wirkung"], WIRKUNG_TEMPLATE.format(erschwerung=erschwerung))
        updated.append((row, reference, erschwerung))

    workbook.save(WERTE)
    print(f"updated={len(updated)}")
    for row, reference, erschwerung in updated:
        print(row, reference, erschwerung)


if __name__ == "__main__":
    main()
