"""LLM-023 Nachtrag (Owner-Entscheidung 2026-09-29): auch die vier zuvor ausgenommenen Zauber
(Wiederbelebung, Eisenhaut, Korpi ignorieren, Durchblick) auf die grad-basierte Kurve normalisieren.
Wiederbelebung ist NPC-only, daher kein Grund fuer eine Ausnahme.
"""
import sys
from pathlib import Path

ROOT = Path(r"E:\Das Western Rollenspiel\LLM")
sys.path.insert(0, str(ROOT / ".python-deps"))
sys.path.insert(0, str(ROOT / "scripts"))
from openpyxl import load_workbook
from backup_werte import backup
from apply_llm_023_stufen_kurve import CURVE, EXCLUDED

WERTE = ROOT / "werte 0.8-claude.xlsx"


def main():
    wb = load_workbook(WERTE)
    werte = wb["Werte"]
    wh = {c.value: c.column for c in werte[1] if c.value}
    grad = {werte.cell(r, wh["Referenz"]).value: int(werte.cell(r, wh["Grad"]).value)
            for r in range(2, werte.max_row + 1) if werte.cell(r, wh["Referenz"]).value in EXCLUDED}
    ws = wb["Spruchmagie-Details"]
    h = {c.value: c.column for c in ws[1] if c.value}
    cols = [h["Stufe1"], h["Stufe2"], h["Stufe3"]]
    todo = [(r, ws.cell(r, h["Referenz"]).value) for r in range(2, ws.max_row + 1)
            if ws.cell(r, h["Referenz"]).value in EXCLUDED]
    if len(todo) != 4 or len(grad) != 4:
        raise SystemExit(f"Erwartet 4 Zauber, gefunden {len(todo)}/{len(grad)}")
    print("Backup:", backup(str(WERTE)))
    for r, ref in todo:
        cur = "/".join(str(ws.cell(r, c).value) for c in cols)
        for c, v in zip(cols, CURVE[grad[ref]]):
            ws.cell(r, c, v)
        print(f"Z.{r} {ref} (Grad {grad[ref]}): {cur} -> {'/'.join(CURVE[grad[ref]])}")
    wb.save(WERTE)


if __name__ == "__main__":
    main()
